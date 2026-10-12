"""Builds outputs/pivot_analysis.xlsx from dataset/netflix_titles.csv.
Run from the Project-1 folder:  python Notes/build_excel.py
Every pivot table is a live COUNTIFS/AVERAGEIFS formula on the Data sheet."""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

CSV = "dataset/netflix_titles.csv"
OUT = "outputs/pivot_analysis.xlsx"
assert os.path.exists(CSV), f"Missing {CSV}. Run from the Project-1 folder."
os.makedirs("outputs", exist_ok=True)

# ---------- clean (same rules as the analysis script) ----------
df = pd.read_csv(CSV).drop_duplicates()
df["year_added"] = pd.to_datetime(df["date_added"].str.strip(), errors="coerce", format="mixed").dt.year
df["country"] = df["country"].fillna("Unknown")
df["primary_country"] = df["country"].str.split(",").str[0].str.strip()
df["minutes"] = df.loc[df["type"] == "Movie", "duration"].str.extract(r"(\d+)", expand=False).astype(float)
df["seasons"] = df.loc[df["type"] == "TV Show", "duration"].str.extract(r"(\d+)", expand=False).astype(float)
bad = df["rating"].astype(str).str.contains("min", na=False)
df.loc[bad, "rating"] = "Unknown"
df["rating"] = df["rating"].fillna("Unknown")
df["genres_padded"] = ", " + df["listed_in"].fillna("") + ","   # lets COUNTIF match exact genre names

cols = ["show_id", "type", "title", "primary_country", "year_added", "release_year",
        "rating", "minutes", "seasons", "listed_in", "genres_padded"]
data = df[cols].copy()
n = len(data)
last = n + 1  # last data row in Excel

# ---------- styles ----------
F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", start_color="1F3A5F")
body = Font(name=F)
bold = Font(name=F, bold=True)
title_font = Font(name=F, bold=True, size=14)
note = Font(name=F, italic=True, color="666666", size=9)
thin = Side(style="thin", color="BBBBBB")

def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

def rng(col_letter):
    return f"Data!${col_letter}$2:${col_letter}${last}"

TYPE, CTRY, YEAR, REL, RATE, MIN, SEAS, PAD = rng("B"), rng("D"), rng("E"), rng("F"), rng("G"), rng("H"), rng("I"), rng("K")

wb = Workbook()

# ---------- Summary ----------
ws = wb.active
ws.title = "Summary"
ws["A1"] = "Netflix Content Analysis: Pivot Workbook"; ws["A1"].font = title_font
ws["A2"] = "All tables are live formulas on the Data sheet. Source: Kaggle 'Netflix Movies and TV Shows'."; ws["A2"].font = note
header(ws, 4, ["Metric", "Value"])
metrics = [
    ("Total titles", f"=COUNTA({rng('A')})", "#,##0"),
    ("Movies", f'=COUNTIF({TYPE},"Movie")', "#,##0"),
    ("TV shows", f'=COUNTIF({TYPE},"TV Show")', "#,##0"),
    ("Movie share of catalog", "=B6/B5", "0.0%"),
    ("TV shows with exactly 1 season (share)", f'=COUNTIFS({TYPE},"TV Show",{SEAS},1)/B7', "0.0%"),
    ("Average movie length (minutes)", f'=AVERAGEIFS({MIN},{TYPE},"Movie")', "0.0"),
    ("TV-MA + TV-14 titles (share)", f'=(COUNTIF({RATE},"TV-MA")+COUNTIF({RATE},"TV-14"))/B5', "0.0%"),
    ("Titles with unknown country (share)", f'=COUNTIF({CTRY},"Unknown")/B5', "0.0%"),
]
for i, (label, f, fmt) in enumerate(metrics, start=5):
    ws.cell(row=i, column=1, value=label).font = body
    c = ws.cell(row=i, column=2, value=f); c.font = body; c.number_format = fmt
ws["A14"] = "Sheets: Pivot_Year_Type, Pivot_Genre, Pivot_Country, Pivot_Rating. Edit the Data sheet and everything recalculates."
ws["A14"].font = note
ws["A15"] = "Genre/country labels on the pivot sheets are the top categories at build time; rerun build_excel.py to refresh them."
ws["A15"].font = note
widths(ws, [46, 16])

# ---------- Data ----------
wd = wb.create_sheet("Data")
header(wd, 1, cols)
for r in data.itertuples(index=False):
    wd.append([None if pd.isna(v) else v for v in r])
for row in wd.iter_rows(min_row=2):
    for c in row:
        c.font = body
wd.freeze_panes = "A2"
wd.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
widths(wd, [10, 10, 38, 18, 11, 12, 10, 9, 9, 40, 40])

# ---------- Pivot: year x type ----------
wy = wb.create_sheet("Pivot_Year_Type")
wy["A1"] = "Titles added per year, by type"; wy["A1"].font = title_font
header(wy, 3, ["Year added", "Movie", "TV Show", "Total", "TV share"])
years = [int(y) for y in sorted(data["year_added"].dropna().unique()) if y >= 2012]
for i, y in enumerate(years, start=4):
    wy.cell(row=i, column=1, value=y).font = body
    wy.cell(row=i, column=2, value=f'=COUNTIFS({YEAR},$A{i},{TYPE},"Movie")')
    wy.cell(row=i, column=3, value=f'=COUNTIFS({YEAR},$A{i},{TYPE},"TV Show")')
    wy.cell(row=i, column=4, value=f"=B{i}+C{i}")
    wy.cell(row=i, column=5, value=f"=IF(D{i}=0,0,C{i}/D{i})").number_format = "0.0%"
    for c in range(2, 6):
        wy.cell(row=i, column=c).font = body
ye = 3 + len(years)
tr = ye + 1
wy.cell(row=tr, column=1, value="Total").font = bold
for c in "BCD":
    wy[f"{c}{tr}"] = f"=SUM({c}4:{c}{ye})"; wy[f"{c}{tr}"].font = bold
wy[f"E{tr}"] = f"=IF(D{tr}=0,0,C{tr}/D{tr})"; wy[f"E{tr}"].number_format = "0.0%"; wy[f"E{tr}"].font = bold
ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Titles added per year"; ch.y_axis.title = "Titles"; ch.x_axis.title = "Year added"
ch.add_data(Reference(wy, min_col=2, max_col=3, min_row=3, max_row=ye), titles_from_data=True)
ch.set_categories(Reference(wy, min_col=1, min_row=4, max_row=ye))
ch.height, ch.width = 9, 18
ch.x_axis.delete = False; ch.y_axis.delete = False
wy.add_chart(ch, "G3")
widths(wy, [14, 10, 10, 10, 10])

# ---------- Pivot: genre ----------
wg = wb.create_sheet("Pivot_Genre")
wg["A1"] = "Top 10 genres (a title with 2 genres counts once in each)"; wg["A1"].font = title_font
header(wg, 3, ["Genre", "Titles", "Share of all titles"])
genre_counts = df["listed_in"].str.split(", ").explode().value_counts().head(10)
for i, g in enumerate(genre_counts.index, start=4):
    wg.cell(row=i, column=1, value=g).font = body
    wg.cell(row=i, column=2, value=f'=COUNTIF({PAD},"*, "&$A{i}&",*")').font = body
    c = wg.cell(row=i, column=3, value=f"=B{i}/Summary!$B$5"); c.number_format = "0.0%"; c.font = body
ch = BarChart(); ch.type = "bar"; ch.title = "Top 10 genres"; ch.legend = None
ch.add_data(Reference(wg, min_col=2, min_row=3, max_row=13), titles_from_data=True)
ch.set_categories(Reference(wg, min_col=1, min_row=4, max_row=13))
ch.height, ch.width = 9, 18; ch.x_axis.scaling.orientation = "maxMin"
ch.x_axis.delete = False; ch.y_axis.delete = False
wg.add_chart(ch, "E3")
widths(wg, [28, 10, 18])

# ---------- Pivot: country ----------
wc = wb.create_sheet("Pivot_Country")
wc["A1"] = "Top 10 primary countries (first listed country; Unknown excluded)"; wc["A1"].font = title_font
header(wc, 3, ["Country", "Titles", "Share of all titles"])
top_c = df[df["primary_country"] != "Unknown"]["primary_country"].value_counts().head(10)
for i, cname in enumerate(top_c.index, start=4):
    wc.cell(row=i, column=1, value=cname).font = body
    wc.cell(row=i, column=2, value=f"=COUNTIF({CTRY},$A{i})").font = body
    c = wc.cell(row=i, column=3, value=f"=B{i}/Summary!$B$5"); c.number_format = "0.0%"; c.font = body
ch = BarChart(); ch.type = "bar"; ch.title = "Top 10 countries"; ch.legend = None
ch.add_data(Reference(wc, min_col=2, min_row=3, max_row=13), titles_from_data=True)
ch.set_categories(Reference(wc, min_col=1, min_row=4, max_row=13))
ch.height, ch.width = 9, 18; ch.x_axis.scaling.orientation = "maxMin"
ch.x_axis.delete = False; ch.y_axis.delete = False
wc.add_chart(ch, "E3")
widths(wc, [22, 10, 18])

# ---------- Pivot: rating x type ----------
wr = wb.create_sheet("Pivot_Rating")
wr["A1"] = "Maturity rating by type (rating = audience maturity label, not a quality score)"; wr["A1"].font = title_font
header(wr, 3, ["Rating", "Movie", "TV Show", "Total", "Share of all titles"])
order = ["TV-Y", "TV-Y7", "TV-G", "G", "TV-PG", "PG", "PG-13", "TV-14", "R", "TV-MA", "NR", "UR", "Unknown"]
present = [r for r in order if r in set(df["rating"])]
for i, rt in enumerate(present, start=4):
    wr.cell(row=i, column=1, value=rt).font = body
    wr.cell(row=i, column=2, value=f'=COUNTIFS({RATE},$A{i},{TYPE},"Movie")').font = body
    wr.cell(row=i, column=3, value=f'=COUNTIFS({RATE},$A{i},{TYPE},"TV Show")').font = body
    wr.cell(row=i, column=4, value=f"=B{i}+C{i}").font = body
    c = wr.cell(row=i, column=5, value=f"=D{i}/Summary!$B$5"); c.number_format = "0.0%"; c.font = body
re_ = 3 + len(present)
ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Maturity rating mix"
ch.add_data(Reference(wr, min_col=2, max_col=3, min_row=3, max_row=re_), titles_from_data=True)
ch.set_categories(Reference(wr, min_col=1, min_row=4, max_row=re_))
ch.height, ch.width = 9, 18; ch.x_axis.delete = False; ch.y_axis.delete = False
wr.add_chart(ch, "G3")
widths(wr, [12, 10, 10, 10, 18])

wb.save(OUT)
print(f"Saved {OUT} ({n} titles). Open it in Excel.")
